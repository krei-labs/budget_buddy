from flask import Blueprint, render_template, request, redirect, url_for, send_file, flash, current_app
from flask_login import login_required, current_user
from app import db
from app.models import Transaction, User, Ticket  # Added Ticket
from datetime import datetime, timedelta, date  # Added date and timedelta for filtering
import io
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from werkzeug.security import check_password_hash

main = Blueprint('main', __name__)

CATEGORIES = ['Food', 'Transport', 'Bills', 'Salary', 'Entertainment', 'Other']

@main.route('/', methods=['GET', 'POST'])
@login_required
def dashboard():
    if request.method == 'POST':
        date_str = request.form.get('date', '').strip()
        category = request.form.get('category', '').strip()
        amount_str = request.form.get('amount', '').strip()
        type_ = request.form.get('type', '').strip()
        notes = request.form.get('notes', '').strip()

        # Validate inputs
        if not date_str or not category or not amount_str or not type_:
            flash('Date, category, amount, and type are required.', 'danger')
            return redirect(url_for('main.dashboard'))

        try:
            date = datetime.strptime(date_str, '%Y-%m-%d').date()
        except ValueError:
            flash('Invalid date format. Please select a valid date.', 'danger')
            return redirect(url_for('main.dashboard'))

        try:
            amount = float(amount_str)
            if amount <= 0:
                flash('Amount must be a positive number.', 'danger')
                return redirect(url_for('main.dashboard'))
        except ValueError:
            flash('Amount must be a valid number.', 'danger')
            return redirect(url_for('main.dashboard'))

        transaction = Transaction(date=date, category=category, amount=amount, type=type_, notes=notes, user_id=current_user.id)
        db.session.add(transaction)
        db.session.commit()
        flash('Transaction added successfully!', 'success')
        return redirect(url_for('main.dashboard'))

    # Get page number from query params (default to 1)
    page = request.args.get('page', 1, type=int)
    per_page = 10  # Change to 20 if you want more rows per page

    # Base query for transactions
    query = Transaction.query.filter_by(user_id=current_user.id)

    # Apply date filter if provided
    date_filter = request.args.get('date_filter', 'all')
    if date_filter == 'today':
        query = query.filter(Transaction.date == datetime.today().date())
    elif date_filter == 'week':
        start_of_week = datetime.today().date() - timedelta(days=datetime.today().weekday())
        query = query.filter(Transaction.date >= start_of_week)
    elif date_filter == 'month':
        start_of_month = datetime.today().date().replace(day=1)
        query = query.filter(Transaction.date >= start_of_month)
    elif date_filter == 'specific' and request.args.get('specific_date'):
        specific_date = datetime.strptime(request.args.get('specific_date'), '%Y-%m-%d').date()
        query = query.filter(Transaction.date == specific_date)

    # Paginate the query
    transactions_paginated = query.paginate(page=page, per_page=per_page, error_out=False)

    # Calculate totals from the full query (not paginated) for accuracy
    all_transactions = query.all()
    total_income = sum(t.amount for t in all_transactions if t.type == 'income')
    total_expenses = sum(t.amount for t in all_transactions if t.type == 'expense')
    balance = total_income - total_expenses

    # Chart data (from all transactions for full breakdown)
    category_totals = {}
    for t in all_transactions:
        if t.type == 'expense':
            category_totals[t.category] = category_totals.get(t.category, 0) + t.amount
    chart_labels = list(category_totals.keys())
    chart_data = list(category_totals.values())

    # Line chart data: Expenses over time (group by date)
    expense_trends = {}
    for t in all_transactions:
        if t.type == 'expense':
            date_str = str(t.date)
            expense_trends[date_str] = expense_trends.get(date_str, 0) + t.amount
    line_labels = sorted(expense_trends.keys())
    line_data = [expense_trends[date] for date in line_labels]

    return render_template('dashboard.html', transactions=transactions_paginated, total_income=total_income,
                           total_expenses=total_expenses, balance=balance, categories=CATEGORIES,
                           chart_labels=chart_labels, chart_data=chart_data,
                           line_labels=line_labels, line_data=line_data, date_filter=date_filter)

@main.route('/delete/<int:id>')
@login_required
def delete_transaction(id):
    transaction = Transaction.query.filter_by(id=id, user_id=current_user.id).first_or_404()
    db.session.delete(transaction)
    db.session.commit()
    flash('Transaction deleted successfully!', 'success')
    return redirect(url_for('main.dashboard'))

@main.route('/export')
@login_required
def export_transactions():
    # Get export filter and dates from request args
    export_filter = request.args.get('export_filter', 'all')
    start_date_str = request.args.get('start_date', '').strip()
    end_date_str = request.args.get('end_date', '').strip()
    
    # Base query
    query = Transaction.query.filter_by(user_id=current_user.id)
    
    # Apply filter based on export option
    if export_filter == 'today':
        query = query.filter(Transaction.date == date.today())
    elif export_filter == 'week':
        week_start = date.today() - timedelta(days=date.today().weekday())
        query = query.filter(Transaction.date >= week_start)
    elif export_filter == 'month':
        month_start = date.today().replace(day=1)
        query = query.filter(Transaction.date >= month_start)
    elif export_filter == 'last_month':
        last_month_end = date.today().replace(day=1) - timedelta(days=1)
        last_month_start = last_month_end.replace(day=1)
        query = query.filter(Transaction.date >= last_month_start, Transaction.date <= last_month_end)
    elif export_filter == 'last_week':
        last_week_end = date.today() - timedelta(days=date.today().weekday())
        last_week_start = last_week_end - timedelta(days=7)
        query = query.filter(Transaction.date >= last_week_start, Transaction.date < last_week_end)
    elif export_filter == 'custom' and start_date_str and end_date_str:
        try:
            start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date()
            end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date()
            query = query.filter(Transaction.date >= start_date, Transaction.date <= end_date)
        except ValueError:
            flash('Invalid date format for custom range.', 'danger')
            return redirect(url_for('main.dashboard'))
    # 'all' does nothing, exports everything
    
    transactions = query.all()  # Get all matching transactions
    
    # Calculate totals for the exported data
    total_income = sum(t.amount for t in transactions if t.type == 'income')
    total_expenses = sum(t.amount for t in transactions if t.type == 'expense')
    balance = total_income - total_expenses
    
    # Generate Word document
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)

    heading = doc.add_heading('Budget Buddy Transactions', 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.runs[0]
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_paragraph(f"Name: {current_user.name}")
    doc.add_paragraph(f"Export Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph(f"Filter: {export_filter.replace('_', ' ').title()}")
    doc.add_paragraph(f"Total Income: ₱{'{:,.2f}'.format(total_income)}")
    doc.add_paragraph(f"Total Expenses: ₱{'{:,.2f}'.format(total_expenses)}")
    doc.add_paragraph(f"Balance: ₱{'{:,.2f}'.format(balance)}")

    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    headers = ['Date', 'Category', 'Amount', 'Type', 'Notes']
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = hdr_cells[i].paragraphs[0].runs[0]
        run.bold = True

    for t in transactions:
        row_cells = table.add_row().cells
        data = [str(t.date), t.category, f"₱{'{:,.2f}'.format(t.amount)}", t.type, t.notes or '']
        for i, cell_data in enumerate(data):
            row_cells[i].text = cell_data
            row_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name='transactions.docx', mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

@main.route('/update_profile', methods=['POST'])
@login_required
def update_profile():
    current_password = request.form.get('current_password', '').strip()
    new_password = request.form.get('new_password', '').strip()
    confirm_password = request.form.get('confirm_password', '').strip()

    # Update profile fields
    current_user.name = request.form['name']
    current_user.age = int(request.form['age'])
    current_user.birthday = datetime.strptime(request.form['birthday'], '%Y-%m-%d').date()
    current_user.email = request.form['email']

    # Handle password change
    if new_password:
        if not check_password_hash(current_user.password_hash, current_password):
            flash('Current password is incorrect.', 'danger')
            return redirect(url_for('main.dashboard'))
        if new_password != confirm_password:
            flash('New passwords do not match.', 'danger')
            return redirect(url_for('main.dashboard'))
        if len(new_password) < 6:
            flash('New password must be at least 6 characters.', 'danger')
            return redirect(url_for('main.dashboard'))
        current_user.set_password(new_password)
        flash('Password updated successfully!', 'success')

    db.session.commit()
    flash('Profile updated successfully!', 'success')
    return redirect(url_for('main.dashboard'))

@main.route('/submit_ticket', methods=['POST'])
@login_required
def submit_ticket():
    message = request.form.get('ticket_message', '').strip()
    if not message:
        flash('Ticket message cannot be empty.', 'danger')
        return redirect(url_for('main.dashboard'))
    ticket = Ticket(user_id=current_user.id, message=message)
    db.session.add(ticket)
    db.session.commit()
    flash('Ticket submitted successfully! Admin will review it.', 'success')
    return redirect(url_for('main.dashboard'))

@main.route('/admin', methods=['GET', 'POST'])
@login_required
def admin():
    if current_user.role != 'admin':
        flash('Access denied.', 'danger')
        return redirect(url_for('main.dashboard'))

    if request.method == 'POST':
        if 'toggle_maintenance' in request.form:
            current_app.config['MAINTENANCE_MODE'] = not current_app.config.get('MAINTENANCE_MODE', False)
            flash('Maintenance mode toggled.', 'success')
        elif 'reset_password' in request.form:
            user_id = request.form.get('user_id')
            user = User.query.get(user_id)
            if user:
                user.set_password('reset123')  # Insecure; use email reset in production
                db.session.commit()
                flash(f'Password reset for {user.username}.', 'success')
        elif 'delete_user' in request.form:
            user_id = request.form.get('user_id')
            user = User.query.get(user_id)
            if user and user.id != current_user.id:  # Prevent self-deletion
                Transaction.query.filter_by(user_id=user_id).delete()  # Delete user's transactions
                db.session.delete(user)
                db.session.commit()
                flash(f'User {user.username} deleted.', 'success')
            else:
                flash('Cannot delete yourself or invalid user.', 'danger')
        elif 'toggle_ban' in request.form:
            user_id = request.form.get('user_id')
            user = User.query.get(user_id)
            if user:
                user.banned = not user.banned
                db.session.commit()
                status = 'banned' if user.banned else 'unbanned'
                flash(f'User {user.username} {status}.', 'success')
        elif 'promote_user' in request.form:
            user_id = request.form.get('user_id')
            user = User.query.get(user_id)
            if user:
                user.role = 'admin'
                db.session.commit()
                flash(f'User {user.username} promoted to admin.', 'success')
        elif 'demote_user' in request.form:
            user_id = request.form.get('user_id')
            user = User.query.get(user_id)
            if user and user.id != current_user.id:  # Prevent self-demotion
                user.role = 'user'
                db.session.commit()
                flash(f'User {user.username} demoted to user.', 'success')
        elif 'resolve_ticket' in request.form:
            ticket_id = request.form.get('ticket_id')
            ticket = Ticket.query.get(ticket_id)
            if ticket:
                ticket.status = 'resolved'
                db.session.commit()
                flash('Ticket resolved.', 'success')
        elif 'create_account' in request.form:
            username = request.form.get('username', '').strip()
            password = request.form.get('password', '').strip()
            email = request.form.get('email', '').strip()  # Added email
            role = request.form.get('role', 'user').strip()

            if not username or not password or not email:  # Added email check
                flash('Username, password, and email are required.', 'danger')
                return redirect(url_for('main.admin'))

            if User.query.filter_by(username=username).first():
                flash('Username already exists.', 'danger')
                return redirect(url_for('main.admin'))

            if User.query.filter_by(email=email).first():
                flash('Email already exists.', 'danger')
                return redirect(url_for('main.admin'))

            user = User(username=username, name='', age=0, birthday=datetime.today().date(), email=email, role=role)  # Set email
            user.set_password(password)
            db.session.add(user)
            db.session.commit()
            flash(f'Account created for {username}. Provide them with the username and password.', 'success')

    users = User.query.all()
    total_users = len(users)
    total_transactions = Transaction.query.count()
    maintenance_mode = current_app.config.get('MAINTENANCE_MODE', False)
    tickets = Ticket.query.filter_by(status='pending').all()  # Pending tickets for review

    return render_template('admin.html', users=users, total_users=total_users, total_transactions=total_transactions, maintenance_mode=maintenance_mode, tickets=tickets)
